import os
import re
import json
import zipfile
from datetime import datetime
import openpyxl

def sanitizar_nome(texto):
    """Remove caracteres inválidos para nomes de pastas e arquivos."""
    if not texto:
        return "OUTROS"
    texto_limpo = re.sub(r'[\\/*?:"<>|]', "", str(texto)).strip()
    return texto_limpo if texto_limpo else "OUTROS"

def formatar_moeda(valor_float):
    """Garante que a Capa tem o visual bonito da moeda brasileira"""
    return f"R$ {valor_float:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

def obter_regras_fornecedor(nome_fornecedor, valor_float):
    nome_upper = str(nome_fornecedor).upper()
    
    regra = {
        "nome_oficial": nome_fornecedor,
        "descricao": "Serviços prestados conforme nota fiscal.",
        "contrato": "",
        "boleto": "sim",
        "valor_igual": "sim",
        "contrato_vigente": "sim",
        "tem_rateio": False,
        "template_file": "Template_PADRAO.xlsx",
        "rateio_coords": {},
        "pasta": sanitizar_nome(nome_fornecedor)
    }
    
    if "CLARO" in nome_upper:
        regra["nome_oficial"] = "Claro S/A"
        regra["pasta"] = "CLARO"
        regra["tem_rateio"] = True
        regra["rateio_coords"] = {
            "fornecedor": "B6", "nf": "B7", "emissao": "F7", 
            "valor": "H7", "periodo": "B8"
        }
        if valor_float <= 500:
            regra["descricao"] = "Conta referente ao pacote Claro Internet Empresa."
            regra["template_file"] = "Template_CLARO_MODENS.xlsx"
        else:
            regra["descricao"] = "Conta telefônica de 29 celulares corporativos."
            regra["template_file"] = "Template_CLARO_CELULARES.xlsx"
            
    elif "VIVO" in nome_upper or "TELEFÔNICA" in nome_upper or "TELEFONICA" in nome_upper:
        regra["nome_oficial"] = "Telefônica Brasil S.A."
        regra["pasta"] = "VIVO"
        regra["descricao"] = "Conta telefônica do terminal 13 3228-3000."
        regra["valor_igual"] = ""
        regra["contrato_vigente"] = ""
        regra["tem_rateio"] = True
        regra["template_file"] = "Template_VIVO.xlsx"
        regra["rateio_coords"] = {
            "fornecedor": "B6", "nf": "B7", "emissao": "F7", 
            "valor": "H7", "periodo": "B8"
        }
        
    elif "INFOCOPPY" in nome_upper or "SOUZA" in nome_upper or "SANTANA" in nome_upper:
        regra["nome_oficial"] = "Souza & Santana Suprimentos e Soluções Técnicas Ltda – ME"
        regra["pasta"] = "INFOCOPPY"
        regra["descricao"] = "48 Impressoras/ Multifuncionais Brother, 5 Impressoras Coloridas Epson, 31 Impressoras Zebras, 10 Multifuncionais Lexmark, 1 Multifuncional Elgin, 2 Scanners Espon/Avision."
        regra["contrato"] = "S052/2022"
        regra["contrato_vigente"] = "nao"
        regra["tem_rateio"] = True
        regra["template_file"] = "Template_INFOCOPPY.xlsx"
        regra["rateio_coords"] = {
            "fornecedor": "B4", "nf": "B5", "emissao": "F5", 
            "valor": "H5", "periodo": "B6"
        }
        
    elif "PCTEC" in nome_upper:
        regra["nome_oficial"] = "PCTEC OUTSOURCING LTDA"
        regra["pasta"] = "PCTEC"
        regra["descricao"] = "Locação de notebooks"
        regra["boleto"] = "nao"
        regra["valor_igual"] = ""
        regra["contrato_vigente"] = ""
        regra["tem_rateio"] = True
        regra["template_file"] = "Template_PCTEC.xlsx"
        regra["rateio_coords"] = {
            "fornecedor": "B6", "nf": "B7", "emissao": "F7", 
            "valor": "H7", "periodo": "F28"
        }

    elif "WTT" in nome_upper:
        regra["nome_oficial"] = "WTT TECNOLOGIA E CONSULTORIA"
        regra["pasta"] = "WTT"
        regra["descricao"] = "Licença de uso, suporte e manutenção do software D-Server, WTT Print e CAP to PACS."
        regra["contrato"] = "S059/2022"
        regra["contrato_vigente"] = "nao"
        
    elif "TASY" in nome_upper or "TUCANO" in nome_upper:
        regra["nome_oficial"] = "TUCANO DO BRASIL SISTEMAS DE INFORMACAO LTDA"
        regra["pasta"] = "TASY"
        regra["descricao"] = "Licença de uso, suporte e manutenção do software Tasy."
        
    return regra

def formatar_checkbox(status):
    if status == "sim": return "(X) Sim", "( ) Não"
    elif status == "nao": return "( ) Sim", "(X) Não"
    else: return "( ) Sim", "( ) Não"

def limpar_valor(valor_str):
    if not valor_str: return 0.0
    v = str(valor_str).upper().replace('R$', '').strip()
    v = v.replace('.', '').replace(',', '.') 
    try: return float(v)
    except ValueError: return 0.0

def encontrar_linha_assinatura(ws):
    for linha in range(1, 150):
        valor_celula = str(ws[f'A{linha}'].value).strip()
        if "Elaborado por:" in valor_celula:
            return linha
    return 69

def _ancora(ws, coordenada):
    for intervalo in ws.merged_cells.ranges:
        if coordenada in intervalo:
            return ws.cell(intervalo.min_row, intervalo.min_col).coordinate
    return coordenada

def _mapa_abas(caminho_xlsx):
    with zipfile.ZipFile(caminho_xlsx) as arquivo:
        relacoes = arquivo.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        livro = arquivo.read("xl/workbook.xml").decode("utf-8")
    destino_por_id = {}
    for encontrado in re.finditer(r"<Relationship\b([^>]*)/>", relacoes):
        atributos = dict(re.findall(r'([\w:]+)="([^"]*)"', encontrado.group(1)))
        destino = atributos.get("Target", "")
        if destino.startswith("worksheets/"):
            destino_por_id[atributos["Id"]] = "xl/" + destino
    abas = {}
    for encontrado in re.finditer(r"<sheet\b([^>]*)/?>", livro):
        atributos = dict(re.findall(r'([\w:]+)="([^"]*)"', encontrado.group(1)))
        abas[atributos["name"]] = destino_por_id[atributos["r:id"]]
    return abas

def _escapar_xml(texto):
    limpo = "".join(caractere for caractere in str(texto) if caractere in "\t\n\r" or ord(caractere) >= 32)
    return limpo.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def _xml_da_celula(referencia, estilo, valor):
    atributo_estilo = f' s="{estilo}"' if estilo else ""
    if isinstance(valor, float):
        return f'<c r="{referencia}"{atributo_estilo}><v>{valor:.2f}</v></c>'
    texto = _escapar_xml("" if valor is None else valor)
    espaco = ' xml:space="preserve"' if texto[:1].isspace() or texto[-1:].isspace() else ""
    return f'<c r="{referencia}"{atributo_estilo} t="inlineStr"><is><t{espaco}>{texto}</t></is></c>'

def _substituir_celula(xml, referencia, valor):
    padrao = re.compile(
        r'<c r="' + re.escape(referencia) + r'"([^>]*?)(?:/>|>.*?</c>)',
        re.S,
    )
    encontrado = padrao.search(xml)
    if not encontrado:
        raise ValueError(f"Célula {referencia} não existe na planilha.")
    estilo = re.search(r'\ss="(\d+)"', encontrado.group(1))
    novo = _xml_da_celula(referencia, estilo.group(1) if estilo else None, valor)
    return xml[:encontrado.start()] + novo + xml[encontrado.end():]

def _gravar_preservando_template(caminho_template, caminho_saida, valores_por_aba):
    """Copia o template e troca só o valor das células.

    O openpyxl regrava o arquivo inteiro e derruba desenhos, imagens e
    comentários. O Excel fecha ao abrir esse arquivo.
    """
    with zipfile.ZipFile(caminho_template, "r") as origem:
        folhas = {}
        for nome_aba, valores in valores_por_aba.items():
            xml = origem.read(nome_aba).decode("utf-8")
            for referencia, valor in valores.items():
                xml = _substituir_celula(xml, referencia, valor)
            folhas[nome_aba] = xml.encode("utf-8")

        relacoes = origem.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        tipos = origem.read("[Content_Types].xml").decode("utf-8")
        if "xl/calcChain.xml" in origem.namelist():
            relacoes = re.sub(r'<Relationship\b[^>]*Target="calcChain.xml"[^>]*/>', "", relacoes)
            tipos = re.sub(r'<Override\b[^>]*PartName="/xl/calcChain.xml"[^>]*/>', "", tipos)

        pasta_saida = os.path.dirname(os.path.abspath(caminho_saida))
        if pasta_saida:
            os.makedirs(pasta_saida, exist_ok=True)
        with zipfile.ZipFile(caminho_saida, "w") as destino:
            for item in origem.infolist():
                if item.filename == "xl/calcChain.xml":
                    continue
                if item.filename in folhas:
                    dados = folhas[item.filename]
                elif item.filename == "xl/_rels/workbook.xml.rels":
                    dados = relacoes.encode("utf-8")
                elif item.filename == "[Content_Types].xml":
                    dados = tipos.encode("utf-8")
                else:
                    dados = origem.read(item.filename)
                destino.writestr(item, dados)

def salvar_na_planilha(dados_json_string, pasta_templates, caminho_saida):
    dados = json.loads(dados_json_string)
    valor_str = dados.get("valor_total", "")
    
    valor_float = limpar_valor(valor_str)
    valor_formatado = formatar_moeda(valor_float)
    
    fornecedor_extraido = dados.get("fornecedor_nome", "")
    regras = obter_regras_fornecedor(fornecedor_extraido, valor_float)
    caminho_template = os.path.join(pasta_templates, regras["template_file"])
    
    if not os.path.exists(caminho_template):
        raise FileNotFoundError(f"Ficheiro de template não encontrado: {caminho_template}")

    wb = openpyxl.load_workbook(caminho_template)
    try:
        return _preencher_e_gravar(wb, caminho_template, caminho_saida, dados, regras, valor_float, valor_formatado)
    finally:
        wb.close()

def _preencher_e_gravar(wb, caminho_template, caminho_saida, dados, regras, valor_float, valor_formatado):
    ws_capa = wb.worksheets[0]
    abas = _mapa_abas(caminho_template)
    alteracoes = {}

    def registrar(ws, coordenada, valor):
        caminho_aba = abas[ws.title]
        alteracoes.setdefault(caminho_aba, {})[_ancora(ws, coordenada)] = valor

    agora = datetime.now()
    data_hoje = agora.strftime("%d/%m/%Y")
    meses = {1:"jan", 2:"fev", 3:"mar", 4:"abr", 5:"mai", 6:"jun", 7:"jul", 8:"ago", 9:"set", 10:"out", 11:"nov", 12:"dez"}
    competencia = f"{meses[agora.month]}/{str(agora.year)[-2:]}"

    # Preenchimento da Capa, sem regravar o desenho e os comentários do template
    registrar(ws_capa, 'B7', regras["nome_oficial"])
    registrar(ws_capa, 'G7', regras["contrato"])
    registrar(ws_capa, 'B9', dados.get("numero_nf", "") or "")
    registrar(ws_capa, 'D9', dados.get("data_vencimento", "") or "")
    registrar(ws_capa, 'E9', valor_formatado)
    registrar(ws_capa, 'H9', valor_formatado)
    registrar(ws_capa, 'B18', regras["descricao"])
    registrar(ws_capa, 'J4', data_hoje)
    registrar(ws_capa, 'D31', competencia)
    
    bol_sim, bol_nao = formatar_checkbox(regras["boleto"])
    registrar(ws_capa, 'B22', f"Boleto  {bol_sim}")
    registrar(ws_capa, 'D22', f"{bol_nao}  - Preencher os dados bancários")

    val_sim, val_nao = formatar_checkbox(regras["valor_igual"])
    registrar(ws_capa, 'B33', f"Valor Faturado é Igual ao Contratado?      {val_sim}          {val_nao}")

    vig_sim, vig_nao = formatar_checkbox(regras["contrato_vigente"])
    registrar(ws_capa, 'B34', f"Contrato e/ou Termo Aditivo Vigente?      {vig_sim}          {vig_nao}")

    aba_rateio_existe = any(aba.lower() == "rateio" for aba in wb.sheetnames)

    if regras["tem_rateio"] and aba_rateio_existe:
        for aba in wb.sheetnames:
            if aba.lower() == "rateio":
                ws_rateio = wb[aba]
                break
                
        coords = regras["rateio_coords"]
        
        registrar(ws_rateio, coords["fornecedor"], regras["nome_oficial"])
        registrar(ws_rateio, coords["nf"], dados.get("numero_nf", "") or "")
        
        data_emissao = dados.get("data_emissao") or dados.get("data_vencimento", "") or ""
        registrar(ws_rateio, coords["emissao"], data_emissao)
        registrar(ws_rateio, coords["valor"], valor_float)
        registrar(ws_rateio, coords["periodo"], competencia)
        
        linha_assinatura = encontrar_linha_assinatura(ws_rateio)
        registrar(ws_rateio, f'B{linha_assinatura}', "Alexandre Siqueira Souza Costa")
        registrar(ws_rateio, f'H{linha_assinatura}', data_hoje)

    _gravar_preservando_template(caminho_template, caminho_saida, alteracoes)
    return caminho_saida