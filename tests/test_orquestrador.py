import json
import os
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from src.modules.batch_processor import processar_fatura, processar_lote_faturas
from src.modules.pdf_exporter import converter_excel_para_pdf
from src.modules.sheet_manager import obter_regras_fornecedor
import openpyxl


PASTA_TEMPLATES = os.path.join(ROOT, "data", "templates")


def _fatura(fornecedor, valor="359,96", numero="160995206"):
    return json.dumps({
        "cnpj_emissor": "00.000.000/0001-00",
        "valor_total": valor,
        "data_vencimento": "24/08/2026",
        "data_emissao": "11/08/2026",
        "fornecedor_nome": fornecedor,
        "numero_nf": numero,
        "linha_digitavel": None,
    })


class TestRegrasFornecedor(unittest.TestCase):
    def test_pastas_homologadas(self):
        casos = [
            ("Claro S/A", "CLARO"),
            ("Souza & Santana Suprimentos", "INFOCOPPY"),
            ("INFOCOPPY", "INFOCOPPY"),
            ("Telefônica Brasil S.A.", "VIVO"),
            ("VIVO", "VIVO"),
            ("PCTEC OUTSOURCING LTDA", "PCTEC"),
            ("WTT TECNOLOGIA", "WTT"),
            ("TUCANO DO BRASIL", "TASY"),
            ("TASY", "TASY"),
            ("", "OUTROS"),
            ("Fornecedor Avulso", "Fornecedor Avulso"),
        ]
        for nome, pasta in casos:
            with self.subTest(nome=nome):
                self.assertEqual(obter_regras_fornecedor(nome, 100)["pasta"], pasta)

    def test_claro_escolhe_template_pelo_valor(self):
        self.assertEqual(
            obter_regras_fornecedor("CLARO", 359.96)["template_file"],
            "Template_CLARO_MODENS.xlsx",
        )
        self.assertEqual(
            obter_regras_fornecedor("CLARO", 1612.38)["template_file"],
            "Template_CLARO_CELULARES.xlsx",
        )


class TestPlanilha(unittest.TestCase):
    def test_claro_preenche_capa_e_preserva_assinatura(self):
        with tempfile.TemporaryDirectory() as saida:
            pdf = os.path.join(saida, "claro.pdf")
            open(pdf, "wb").close()
            with patch("src.modules.batch_processor.extrair_dados_pdf", return_value=_fatura("CLARO")), \
                 patch("src.modules.batch_processor.converter_excel_para_pdf", return_value="ok"):
                resultado = processar_fatura(pdf, PASTA_TEMPLATES, saida)

            self.assertEqual(resultado["status"], "SUCESSO")
            self.assertTrue(resultado["excel"].endswith(os.path.join("CLARO", "fatura_NF_160995206.xlsx")))

            wb = openpyxl.load_workbook(resultado["excel"])
            capa = wb.worksheets[0]
            self.assertEqual(capa["B7"].value, "Claro S/A")
            self.assertEqual(capa["B9"].value, "160995206")
            self.assertEqual(capa["E9"].value, "R$ 359,96")
            self.assertEqual(capa["B18"].value, "Conta referente ao pacote Claro Internet Empresa.")
            self.assertEqual(capa["B7"].font.name, "Verdana")
            self.assertEqual(capa["B7"].font.size, 13)
            self.assertTrue(capa["B7"].font.bold)

            rateio = wb["Rateio"]
            self.assertEqual(rateio["B6"].value, "Claro S/A")
            self.assertEqual(rateio["E7"].value, "11/08/2026")
            self.assertEqual(rateio["H7"].value, 359.96)
            self.assertEqual(rateio["B27"].value, "Alexandre Siqueira Souza Costa")
            self.assertEqual(rateio["B27"].font.name, rateio["A27"].font.name)
            self.assertEqual(rateio["B27"].font.size, rateio["A27"].font.size)
            self.assertRegex(str(rateio["H27"].value), r"\d{2}/\d{2}/\d{4}")

    def test_tasy_usa_template_padrao_sem_aba_de_rateio(self):
        with tempfile.TemporaryDirectory() as saida:
            pdf = os.path.join(saida, "tasy.pdf")
            open(pdf, "wb").close()
            with patch("src.modules.batch_processor.extrair_dados_pdf", return_value=_fatura("TASY", "2186,34", "2286")), \
                 patch("src.modules.batch_processor.converter_excel_para_pdf", return_value="ok"):
                resultado = processar_fatura(pdf, PASTA_TEMPLATES, saida)

            wb = openpyxl.load_workbook(resultado["excel"])
            self.assertEqual(wb.sheetnames, ["CHE"])
            self.assertEqual(wb["CHE"]["B7"].value, "TUCANO DO BRASIL SISTEMAS DE INFORMACAO LTDA")
            self.assertIn("TASY", resultado["excel"])


class TestLote(unittest.TestCase):
    def test_isola_erro_e_grava_relatorio(self):
        with tempfile.TemporaryDirectory() as raiz:
            entrada = os.path.join(raiz, "input")
            saida = os.path.join(raiz, "output")
            os.makedirs(entrada)
            for nome in ("ok.pdf", "ruim.pdf"):
                open(os.path.join(entrada, nome), "wb").close()

            def extrair(caminho):
                if caminho.endswith("ruim.pdf"):
                    raise RuntimeError("leitura falhou")
                return _fatura("VIVO", "263,17", "686186")

            with patch("src.modules.batch_processor.extrair_dados_pdf", side_effect=extrair), \
                 patch("src.modules.batch_processor.converter_excel_para_pdf", return_value="ok"):
                resultados = processar_lote_faturas(entrada, saida, PASTA_TEMPLATES)

            self.assertEqual([r["status"] for r in resultados], ["SUCESSO", "ERRO"])
            self.assertTrue(os.path.exists(os.path.join(saida, "VIVO", "fatura_NF_686186.xlsx")))
            with open(os.path.join(saida, "relatorio_lote.json"), encoding="utf-8") as arquivo:
                relatorio = json.load(arquivo)
            self.assertEqual(relatorio[1]["erro"], "leitura falhou")
            self.assertIsNone(relatorio[1]["excel"])

    def test_pdf_falho_mantem_excel_como_parcial(self):
        with tempfile.TemporaryDirectory() as raiz:
            entrada = os.path.join(raiz, "input")
            saida = os.path.join(raiz, "output")
            os.makedirs(entrada)
            open(os.path.join(entrada, "wtt.pdf"), "wb").close()

            with patch("src.modules.batch_processor.extrair_dados_pdf", return_value=_fatura("WTT", "1000,00", "55")), \
                 patch("src.modules.batch_processor.converter_excel_para_pdf", side_effect=RuntimeError("sem excel")):
                resultados = processar_lote_faturas(entrada, saida, PASTA_TEMPLATES)

            self.assertEqual(resultados[0]["status"], "PARCIAL")
            self.assertTrue(os.path.exists(resultados[0]["excel"]))
            self.assertIsNone(resultados[0]["pdf"])
            self.assertIn("sem excel", resultados[0]["erro"])

    def test_mesma_nf_no_lote_nao_sobrescreve(self):
        with tempfile.TemporaryDirectory() as raiz:
            entrada = os.path.join(raiz, "input")
            saida = os.path.join(raiz, "output")
            os.makedirs(entrada)
            for nome in ("a.pdf", "b.pdf"):
                open(os.path.join(entrada, nome), "wb").close()

            with patch("src.modules.batch_processor.extrair_dados_pdf", return_value=_fatura("PCTEC", "5171,89", "341")), \
                 patch("src.modules.batch_processor.converter_excel_para_pdf", return_value="ok"):
                resultados = processar_lote_faturas(entrada, saida, PASTA_TEMPLATES)

            excels = {os.path.basename(r["excel"]) for r in resultados}
            self.assertEqual(excels, {"fatura_NF_341.xlsx", "fatura_NF_341_b.xlsx"})


class TestPdfExporter(unittest.TestCase):
    def test_arquivo_inexistente(self):
        with self.assertRaises(FileNotFoundError):
            converter_excel_para_pdf("/tmp/nao-existe.xlsx", "/tmp/nao-existe.pdf")

    def test_libreoffice_quando_nao_e_windows(self):
        with tempfile.TemporaryDirectory() as pasta:
            origem = os.path.join(pasta, "fatura.xlsx")
            destino = os.path.join(pasta, "saida.pdf")
            with open(origem, "wb") as arquivo:
                arquivo.write(b"PK")

            def run(cmd, **kwargs):
                outdir = cmd[cmd.index("--outdir") + 1]
                with open(os.path.join(outdir, "fatura.pdf"), "wb") as gerado_tmp:
                    gerado_tmp.write(b"%PDF-1.4")
                return subprocess.CompletedProcess(cmd, 0, b"", b"")

            with patch("src.modules.pdf_exporter.os.name", "posix"), \
                 patch("src.modules.pdf_exporter.shutil.which", return_value="/usr/bin/soffice"), \
                 patch("src.modules.pdf_exporter.subprocess.run", side_effect=run):
                gerado = converter_excel_para_pdf(origem, destino)

            self.assertEqual(gerado, os.path.abspath(destino))
            with open(destino, "rb") as arquivo:
                self.assertEqual(arquivo.read(), b"%PDF-1.4")

    def test_sem_libreoffice_explica_o_motivo(self):
        with tempfile.TemporaryDirectory() as pasta:
            origem = os.path.join(pasta, "fatura.xlsx")
            open(origem, "wb").close()
            with patch("src.modules.pdf_exporter.os.name", "posix"), \
                 patch("src.modules.pdf_exporter.shutil.which", return_value=None):
                with self.assertRaises(RuntimeError) as ctx:
                    converter_excel_para_pdf(origem, os.path.join(pasta, "out.pdf"))
            self.assertIn("LibreOffice", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
